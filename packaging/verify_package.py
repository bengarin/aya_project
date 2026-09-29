"""
verify_package.py
-----------------
Verifie que l'EXE construit donne EXACTEMENT le meme resultat que le code source.

  1. genere les fichiers de test (tests/make_fixtures.py : les 10 cas obligatoires)
  2. lance AYA_Excel_CLI.exe dessus            -> resultat "EXE"
  3. lance le code source (pipeline) dessus    -> resultat "SOURCE"
  4. compare les deux fichiers cellule par cellule (valeurs, formules, couleurs)
     + verifie que le rapport est cree et que les fichiers sources n'ont pas bouge

Usage :
  python packaging/verify_package.py dist/AYA_Excel/AYA_Excel_CLI.exe
  python packaging/verify_package.py wine dist/AYA_Excel/AYA_Excel_CLI.exe   (Linux + Wine)
Code retour 0 = OK.
"""

from __future__ import annotations

import hashlib
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

from openpyxl import load_workbook                  # noqa: E402

from make_fixtures import build_all                 # noqa: E402
from pipeline import Selection, Session             # noqa: E402
from processor import Options                       # noqa: E402


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _to_exe_path(path: Path, use_wine: bool) -> str:
    return "Z:" + str(path).replace("/", "\\") if use_wine else str(path)


def _snapshot(path: Path) -> dict:
    wb = load_workbook(path)
    data = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                fill = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else None
                if cell.value is not None or fill:
                    data[(ws.title, cell.coordinate)] = (cell.value, fill)
    wb.close()
    return data


def main(argv: list[str]) -> int:
    if not argv:
        print(__doc__)
        return 2
    use_wine = argv[0].lower() == "wine"
    command = list(argv)

    with tempfile.TemporaryDirectory() as tmp:
        folder = Path(tmp)
        files = build_all(folder)
        before = {k: _sha(p) for k, p in files.items()}

        out_exe = folder / "resultat_exe.xlsx"
        cmd = command + ["--cli",
                         "-r", _to_exe_path(files["reference"], use_wine),
                         "-t", _to_exe_path(files["target"], use_wine),
                         "-k", _to_exe_path(files["kam"], use_wine),
                         "-o", _to_exe_path(out_exe, use_wine)]
        print("> " + " ".join(cmd))
        proc = subprocess.run(cmd, capture_output=True, text=True, errors="replace", timeout=300)
        print(proc.stdout[-3000:])
        if proc.returncode != 0 or not out_exe.exists():
            print(proc.stderr[-3000:])
            print(f"ECHEC : l'EXE a retourne {proc.returncode}")
            return 1

        out_src = folder / "resultat_source.xlsx"
        session = Session(Selection(reference=files["reference"], target=files["target"],
                                    kam=files["kam"]), Options())
        try:
            session.load()
            session.analyze()
            session.apply(out_src, write_report=False)
        finally:
            session.close()

        errors = []
        if {k: _sha(p) for k, p in files.items()} != before:
            errors.append("un fichier SOURCE a ete modifie (interdit)")
        reports = sorted(p.name for p in folder.glob("resultat_exe*RAPPORT*"))
        if not reports:
            errors.append("aucun rapport genere a cote du resultat")
        a, b = _snapshot(out_exe), _snapshot(out_src)
        diff = [k for k in sorted(set(a) | set(b)) if a.get(k) != b.get(k)]
        if diff:
            errors.append(f"{len(diff)} cellule(s) differentes, ex : "
                          + ", ".join(f"{s}!{c}: exe={a.get((s, c))} source={b.get((s, c))}"
                                      for s, c in diff[:5]))

        print(f"Cellules comparees : {len(a)} | rapports : {reports}")
        if errors:
            for e in errors:
                print("ECHEC : " + e)
            return 1
        print("OK : l'EXE donne exactement le meme resultat que le code source.")
        return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
