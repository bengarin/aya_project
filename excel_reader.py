"""
excel_reader.py
---------------
Lecture des 3 fichiers Excel.

Principe : on ne travaille JAMAIS avec des numeros de colonnes ecrits en dur.
On lit la ligne d'en-tetes, on reconnait les colonnes par leur NOM, et on
memorise "quel nom = quelle colonne". Comme ca, si l'utilisateur deplace une
colonne, le logiciel continue de fonctionner.

Les fichiers 1 (Reference) et 3 (KAM) sont ouverts en LECTURE SEULE.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from matcher import (
    cell_to_text,
    norm_code,
    norm_division,
    norm_store,
    norm_text,
    split_divisions,
)

# Nombre de lignes du haut dans lesquelles on cherche la ligne d'en-tetes.
HEADER_SCAN_ROWS = 15
# On arrete de chercher des donnees apres autant de lignes vides consecutives.
BLANK_STREAK_STOP = 150


class ExcelReadError(Exception):
    """Erreur de lecture claire, affichable telle quelle dans l'interface."""


# ---------------------------------------------------------------------------
# Description des colonnes attendues dans chaque fichier
# ---------------------------------------------------------------------------
@dataclass(frozen=True)
class SheetSpec:
    label: str
    required: tuple[str, ...]
    optional: tuple[str, ...] = ()
    aliases: dict[str, tuple[str, ...]] = field(default_factory=dict)

    @property
    def all_columns(self) -> tuple[str, ...]:
        return self.required + self.optional


REFERENCE_SPEC = SheetSpec(
    label="Fichier 1 - Reference",
    required=("ID Promoter", "Store Code", "PROMOTER", "STORE", "Divisions", "Working Days"),
    optional=("CITY", "GRADE", "Family", "Sub-channel"),
    aliases={
        "ID Promoter": ("id promoteur", "idpromoter", "id promo"),
        "Store Code": ("code store", "code magasin", "storecode"),
        "PROMOTER": ("promoteur", "nom promoteur"),
        "STORE": ("magasin", "store name", "nom store"),
        "Divisions": ("division", "divisions1"),
        "Working Days": ("working day", "day", "days", "jours", "nb jours"),
        "CITY": ("ville",),
    },
)

TARGET_SPEC = SheetSpec(
    label="Fichier 2 - BDD a traiter",
    required=(
        "ID Promoter", "Code Store", "Promoter", "Store",
        "Division1", "KAM", "DAY", "Column1",
    ),
    optional=("Annee", "Mois", "Division", "City", "STATUT", "IDAYA"),
    aliases={
        "ID Promoter": ("id promoteur", "idpromoter"),
        "Code Store": ("store code", "code magasin"),
        "Promoter": ("promoteur",),
        "Store": ("magasin", "store name"),
        "Division1": ("division 1", "division1", "sous division"),
        "DAY": ("days", "working days", "jours"),
        "Column1": ("colonne1", "column 1", "colonne 1", "cle", "key", "id"),
        "Annee": ("annee", "année", "year"),
        "Mois": ("month",),
        "City": ("ville",),
        "STATUT": ("statut", "status"),
        "IDAYA": ("id aya", "idaya"),
    },
)

KAM_SPEC = SheetSpec(
    label="Fichier 3 - Affectation KAM",
    required=("Store", "KAM"),
    optional=("CITY", "Code store"),
    aliases={
        "Store": ("magasin", "store name", "nom store"),
        "KAM": ("kam", "responsable"),
        "CITY": ("city", "ville"),
        "Code store": ("store code", "code magasin"),
    },
)


# ---------------------------------------------------------------------------
# Detection feuille + colonnes
# ---------------------------------------------------------------------------
@dataclass
class SheetLayout:
    """Ce que le logiciel a compris d'une feuille : ou sont les en-tetes et les colonnes."""

    title: str
    header_row: int
    columns: dict[str, int]          # nom canonique -> numero de colonne (1 = A)
    missing: list[str]               # colonnes obligatoires absentes
    last_data_row: int
    visible: bool = True
    table_name: str | None = None
    header_labels: dict[int, str] = field(default_factory=dict)

    @property
    def is_valid(self) -> bool:
        return not self.missing

    @property
    def data_rows(self) -> int:
        return max(0, self.last_data_row - self.header_row)

    def col(self, name: str) -> int | None:
        return self.columns.get(name)

    def describe(self) -> str:
        found = ", ".join(f"{k}=col.{_letter(v)}" for k, v in sorted(self.columns.items(), key=lambda x: x[1]))
        return f"Feuille '{self.title}' | en-tetes ligne {self.header_row} | {self.data_rows} lignes | {found}"


def _letter(index: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(index)


def _match_columns(header_values: dict[int, object], spec: SheetSpec) -> dict[str, int]:
    """Associe chaque colonne attendue a un numero de colonne reel.

    Etape 1 : correspondance EXACTE du nom (apres nettoyage).
    Etape 2 : correspondance par synonyme, seulement pour ce qui manque encore.
    """
    normalized = {idx: norm_text(val) for idx, val in header_values.items() if cell_to_text(val)}
    result: dict[str, int] = {}
    used: set[int] = set()

    for name in spec.all_columns:
        target = norm_text(name)
        for idx, value in normalized.items():
            if idx in used:
                continue
            if value == target:
                result[name] = idx
                used.add(idx)
                break

    for name in spec.all_columns:
        if name in result:
            continue
        candidates = {norm_text(a) for a in spec.aliases.get(name, ())}
        if not candidates:
            continue
        for idx, value in normalized.items():
            if idx in used:
                continue
            if value in candidates:
                result[name] = idx
                used.add(idx)
                break
    return result


def _find_last_data_row(ws, columns: dict[str, int], header_row: int, hard_stop: int) -> int:
    """Trouve la derniere ligne qui contient vraiment des donnees."""
    indexes = sorted(columns.values())
    if not indexes:
        return header_row
    last = header_row
    blank_streak = 0
    row = header_row + 1
    while row <= hard_stop:
        has_value = any(cell_to_text(ws.cell(row, idx).value) for idx in indexes)
        if has_value:
            last = row
            blank_streak = 0
        else:
            blank_streak += 1
            if blank_streak >= BLANK_STREAK_STOP:
                break
        row += 1
    return last


def analyze_sheet(ws, spec: SheetSpec) -> SheetLayout:
    """Analyse UNE feuille : trouve la ligne d'en-tetes puis les colonnes."""
    max_col = min(ws.max_column or 1, 80)
    best: tuple[int, dict[str, int]] | None = None
    for row in range(1, min(ws.max_row or 1, HEADER_SCAN_ROWS) + 1):
        header_values = {c: ws.cell(row, c).value for c in range(1, max_col + 1)}
        columns = _match_columns(header_values, spec)
        score = sum(1 for name in spec.required if name in columns)
        if best is None or score > best[0]:
            best = (score, {"row": row, "columns": columns})
        if score == len(spec.required):
            break

    header_row = best[1]["row"] if best else 1
    columns = best[1]["columns"] if best else {}
    missing = [name for name in spec.required if name not in columns]

    # Une feuille-tableau Excel donne directement la derniere ligne de donnees.
    table_name = None
    table_end = None
    for table in getattr(ws, "tables", {}).values():
        try:
            from openpyxl.utils.cell import range_boundaries

            min_c, min_r, max_c, max_r = range_boundaries(table.ref)
        except Exception:                                   # pragma: no cover - securite
            continue
        if min_r <= header_row <= max_r:
            table_name = table.name
            table_end = max_r
            break

    hard_stop = min(ws.max_row or header_row, header_row + 200_000)
    last_row = _find_last_data_row(ws, columns, header_row, hard_stop) if columns else header_row
    if table_end:
        last_row = max(last_row, table_end)

    header_labels = {c: cell_to_text(ws.cell(header_row, c).value) for c in range(1, max_col + 1)}
    return SheetLayout(
        title=ws.title,
        header_row=header_row,
        columns=columns,
        missing=missing,
        last_data_row=last_row,
        visible=(getattr(ws, "sheet_state", "visible") == "visible"),
        table_name=table_name,
        header_labels={k: v for k, v in header_labels.items() if v},
    )


def analyze_workbook(wb, spec: SheetSpec) -> list[SheetLayout]:
    """Analyse toutes les feuilles et classe les meilleures candidates en premier."""
    layouts = [analyze_sheet(ws, spec) for ws in wb.worksheets]
    layouts.sort(key=lambda l: (not l.is_valid, not l.visible, -l.data_rows, l.title))
    return layouts


def pick_layout(layouts: list[SheetLayout], spec: SheetSpec, wanted: str | None = None) -> SheetLayout:
    """Choisit la feuille a utiliser (celle demandee, sinon la meilleure candidate)."""
    if wanted:
        for layout in layouts:
            if layout.title == wanted:
                if not layout.is_valid:
                    raise ExcelReadError(
                        f"{spec.label} : la feuille '{wanted}' n'a pas les colonnes obligatoires : "
                        + ", ".join(layout.missing)
                    )
                return layout
        raise ExcelReadError(f"{spec.label} : feuille '{wanted}' introuvable.")
    for layout in layouts:
        if layout.is_valid:
            return layout
    detail = layouts[0].missing if layouts else spec.required
    raise ExcelReadError(
        f"{spec.label} : aucune feuille ne contient les colonnes obligatoires.\n"
        f"Colonnes manquantes (meilleure feuille testee) : " + ", ".join(detail)
    )


# ---------------------------------------------------------------------------
# Fichier 1 : Reference (LECTURE SEULE)
# ---------------------------------------------------------------------------
@dataclass
class ReferenceRow:
    source_row: int
    id_promoter: str
    store_code: str
    promoter: str
    store: str
    divisions_raw: str
    divisions: list[str]
    working_days: str
    city: str = ""
    grade: str = ""


@dataclass
class ReferenceData:
    path: Path
    layout: SheetLayout
    rows: list[ReferenceRow]
    by_store: dict[str, list[ReferenceRow]]
    sheet_names: list[str]

    def get(self, store: str) -> list[ReferenceRow]:
        return self.by_store.get(norm_store(store), [])


def load_reference(path: str | Path, sheet: str | None = None) -> ReferenceData:
    path = Path(path)
    wb = load_workbook(path, data_only=True, read_only=False)
    try:
        sheet_names = wb.sheetnames
        layouts = analyze_workbook(wb, REFERENCE_SPEC)
        layout = pick_layout(layouts, REFERENCE_SPEC, sheet)
        ws = wb[layout.title]
        rows: list[ReferenceRow] = []
        by_store: dict[str, list[ReferenceRow]] = {}
        for r in range(layout.header_row + 1, layout.last_data_row + 1):
            def val(name: str) -> str:
                idx = layout.col(name)
                return cell_to_text(ws.cell(r, idx).value) if idx else ""

            store = val("STORE")
            divisions_raw = val("Divisions")
            if not store and not divisions_raw and not val("ID Promoter"):
                continue                                    # ligne totalement vide
            entry = ReferenceRow(
                source_row=r,
                id_promoter=val("ID Promoter"),
                store_code=val("Store Code"),
                promoter=val("PROMOTER"),
                store=store,
                divisions_raw=divisions_raw,
                divisions=split_divisions(divisions_raw),
                working_days=val("Working Days"),
                city=val("CITY"),
                grade=val("GRADE"),
            )
            rows.append(entry)
            by_store.setdefault(norm_store(store), []).append(entry)
        return ReferenceData(path=path, layout=layout, rows=rows, by_store=by_store, sheet_names=sheet_names)
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# Fichier 3 : Affectation KAM (LECTURE SEULE)
# ---------------------------------------------------------------------------
@dataclass
class KamEntry:
    store: str
    kam: str
    store_code: str
    sheet: str
    division: str          # "VD", "DA" ou "" (feuille generique)
    source_row: int


@dataclass
class KamData:
    path: Path
    layouts: list[SheetLayout]
    entries: list[KamEntry]
    sheet_names: list[str]


def load_kam(path: str | Path, sheets: list[str] | None = None) -> KamData:
    """Lit le fichier des KAM.

    Cas reel important : ce fichier contient une feuille 'VD' et une feuille 'DA'.
    Le KAM depend donc du magasin ET de la division. On memorise la division
    d'apres le nom de la feuille. Si la feuille ne s'appelle ni VD ni DA, elle est
    consideree comme "generique" (valable pour toutes les divisions).
    """
    path = Path(path)
    wb = load_workbook(path, data_only=True, read_only=False)
    try:
        sheet_names = wb.sheetnames
        layouts_all = [analyze_sheet(wb[name], KAM_SPEC) for name in sheet_names]
        usable = [l for l in layouts_all if l.is_valid and (sheets is None or l.title in sheets)]
        if not usable:
            missing = layouts_all[0].missing if layouts_all else KAM_SPEC.required
            raise ExcelReadError(
                "Fichier 3 - Affectation KAM : aucune feuille avec les colonnes obligatoires "
                f"({', '.join(KAM_SPEC.required)}). Manquantes : {', '.join(missing)}"
            )
        entries: list[KamEntry] = []
        for layout in usable:
            ws = wb[layout.title]
            division = norm_division(layout.title)
            if division not in ("VD", "DA"):
                division = ""
            for r in range(layout.header_row + 1, layout.last_data_row + 1):
                store = cell_to_text(ws.cell(r, layout.col("Store")).value) if layout.col("Store") else ""
                kam = cell_to_text(ws.cell(r, layout.col("KAM")).value) if layout.col("KAM") else ""
                if not store or not kam:
                    continue
                code_idx = layout.col("Code store")
                entries.append(
                    KamEntry(
                        store=store,
                        kam=kam,
                        store_code=norm_code(ws.cell(r, code_idx).value) if code_idx else "",
                        sheet=layout.title,
                        division=division,
                        source_row=r,
                    )
                )
        return KamData(path=path, layouts=usable, entries=entries, sheet_names=sheet_names)
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# Fichier 2 : BDD a traiter (lu ET modifie - mais jamais a la source)
# ---------------------------------------------------------------------------
@dataclass
class TargetRow:
    row: int
    values: dict[str, str]           # nom canonique -> texte affiche (valeur calculee incluse)

    def get(self, name: str) -> str:
        return self.values.get(name, "")


@dataclass
class TargetData:
    path: Path
    workbook: object                 # classeur openpyxl ouvert en ecriture (formules conservees)
    worksheet: object
    layout: SheetLayout
    rows: list[TargetRow]
    sheet_names: list[str]
    warnings: list[str] = field(default_factory=list)

    def close(self) -> None:
        try:
            self.workbook.close()
        except Exception:                                   # pragma: no cover
            pass


def load_target(path: str | Path, sheet: str | None = None) -> TargetData:
    """Charge le fichier 2.

    On l'ouvre 2 fois :
      1. en mode "formules" (data_only=False) : c'est ce classeur qu'on modifiera
         et qui garde couleurs, tableaux, filtres et formules ;
      2. en mode "valeurs" (data_only=True)   : pour lire le RESULTAT des formules
         (ex: si Column1 est une formule, on a besoin de son resultat pour comparer).
    """
    path = Path(path)
    wb = load_workbook(path, data_only=False, keep_vba=str(path).lower().endswith(".xlsm"))
    warnings: list[str] = []
    try:
        layouts = analyze_workbook(wb, TARGET_SPEC)
        layout = pick_layout(layouts, TARGET_SPEC, sheet)
        ws = wb[layout.title]

        cached: dict[tuple[int, int], object] = {}
        try:
            wbv = load_workbook(path, data_only=True, read_only=True)
            wsv = wbv[layout.title]
            indexes = sorted(layout.columns.values())
            max_idx = max(indexes) if indexes else 1
            for row in wsv.iter_rows(
                min_row=layout.header_row + 1, max_row=layout.last_data_row, max_col=max_idx
            ):
                for cell in row:
                    if cell.value is not None:
                        cached[(cell.row, cell.column)] = cell.value
            wbv.close()
        except Exception as exc:                            # pragma: no cover - fichier exotique
            warnings.append(f"Valeurs calculees non lisibles ({exc}). Les formules seront comparees telles quelles.")

        rows: list[TargetRow] = []
        formula_without_value = 0
        for r in range(layout.header_row + 1, layout.last_data_row + 1):
            values: dict[str, str] = {}
            for name, idx in layout.columns.items():
                raw = ws.cell(r, idx).value
                if isinstance(raw, str) and raw.startswith("="):
                    resolved = cached.get((r, idx))
                    if resolved is None:
                        formula_without_value += 1
                        values[name] = ""
                    else:
                        values[name] = cell_to_text(resolved)
                else:
                    values[name] = cell_to_text(raw if raw is not None else cached.get((r, idx)))
            rows.append(TargetRow(row=r, values=values))
        if formula_without_value:
            warnings.append(
                f"{formula_without_value} cellule(s) contiennent une formule sans resultat enregistre "
                "(le fichier n'a jamais ete recalcule par Excel). Ces cellules sont traitees comme vides."
            )
        return TargetData(
            path=path, workbook=wb, worksheet=ws, layout=layout,
            rows=rows, sheet_names=wb.sheetnames, warnings=warnings,
        )
    except Exception:
        wb.close()
        raise


# ---------------------------------------------------------------------------
# Utilitaires "rapides" (sans ouvrir tout le classeur)
# ---------------------------------------------------------------------------
def list_sheet_names(path: str | Path) -> list[str]:
    """Liste les feuilles d'un .xlsx sans le charger entierement (tres rapide).

    Utilise pour remplir les listes deroulantes de l'interface.
    """
    import re
    import zipfile

    try:
        with zipfile.ZipFile(path) as archive:
            xml = archive.read("xl/workbook.xml").decode("utf-8", "ignore")
    except Exception:                                       # pragma: no cover - fichier illisible
        return []
    names = re.findall(r'<sheet[^>]*name="([^"]+)"', xml)
    unescape = {"&amp;": "&", "&lt;": "<", "&gt;": ">", "&quot;": '"', "&apos;": "'"}
    result = []
    for name in names:
        for key, value in unescape.items():
            name = name.replace(key, value)
        result.append(name)
    return result


# Elements Excel qu'openpyxl ne sait pas recopier dans le fichier resultat.
_UNSUPPORTED_PARTS = (
    ("xl/connections.xml", "connexions de donnees externes"),
    ("queryTables/", "tableaux de requete / Power Query"),
    ("pivotCache/", "tableaux croises dynamiques"),
    ("charts/", "graphiques"),
    ("xl/media/", "images"),
    ("slicer", "segments (slicers)"),
    ("printerSettings", "parametres d'impression"),
    ("namedSheetViews", "vues de feuille nommees"),
)


def detect_unsupported_parts(path: str | Path) -> list[str]:
    """Previent l'utilisateur de ce qui ne sera PAS recopie dans le fichier resultat."""
    import zipfile

    try:
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
    except Exception:                                       # pragma: no cover
        return []
    found = []
    for marker, label in _UNSUPPORTED_PARTS:
        if any(marker in name for name in names) and label not in found:
            found.append(label)
    return found
