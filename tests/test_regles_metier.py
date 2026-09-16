"""
tests/test_regles_metier.py
---------------------------
Verifie que le logiciel respecte les regles metier, une regle = un test.

Lancement :
    python -m unittest discover -s tests -v
    (ou simplement : python tests/test_regles_metier.py)
"""

from __future__ import annotations

import hashlib
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import openpyxl                                              # noqa: E402

from make_fixtures import build_all                          # noqa: E402
from matcher import build_key, norm_store, split_divisions   # noqa: E402
from pipeline import PipelineError, Selection, Session       # noqa: E402
from processor import Options                                # noqa: E402

SHEET = "BDD PROMOTERS MONTH"
COL = {name: index for index, name in enumerate(
    ["Annee", "Mois", "Division", "Division1", "KAM", "ID Promoter", "Code Store",
     "Promoter", "Store", "City", "STATUT", "DAY", "IDAYA", "Column1"], start=1)}


def md5(path: Path) -> str:
    return hashlib.md5(Path(path).read_bytes()).hexdigest()


def cell(ws, row: int, column: str):
    return ws.cell(row, COL[column]).value


class BaseCase(unittest.TestCase):
    """Prepare 3 fichiers de test neufs pour chaque test (isolation totale)."""

    def setUp(self) -> None:
        self.folder = Path(tempfile.mkdtemp(prefix="aya_test_"))
        self.files = build_all(self.folder)
        self.output = self.folder / "RESULTAT.xlsx"
        self.hashes = {name: md5(path) for name, path in self.files.items()}

    def tearDown(self) -> None:
        shutil.rmtree(self.folder, ignore_errors=True)

    def run_pipeline(self, options: Options | None = None, output: Path | None = None):
        session = Session(
            Selection(reference=self.files["reference"], target=self.files["target"], kam=self.files["kam"]),
            options or Options(),
        )
        try:
            session.load()
            plan = session.analyze()
            result, _reports = session.apply(output or self.output, write_report=False)
            return plan, result
        finally:
            session.close()

    def result_sheet(self, path: Path | None = None):
        wb = openpyxl.load_workbook(path or self.output)
        return wb, wb[SHEET]


class TestDetection(BaseCase):
    def test_feuilles_et_colonnes_detectees(self):
        session = Session(Selection(reference=self.files["reference"], target=self.files["target"],
                                    kam=self.files["kam"]))
        try:
            data = session.load()
            # en-tetes en ligne 2 dans la Reference, colonnes reconnues par leur nom
            self.assertEqual(data.reference.layout.header_row, 2)
            self.assertEqual(data.reference.layout.col("STORE"), 6)
            # la feuille parasite 'Notes' ne doit pas etre choisie
            self.assertEqual(data.target.layout.title, SHEET)
            self.assertEqual(data.target.layout.header_row, 1)
            # le fichier 3 expose bien une feuille par division
            self.assertEqual(data.kam.divisions, ["DA", "VD"])
        finally:
            session.close()

    def test_colonne_obligatoire_manquante(self):
        wb = openpyxl.load_workbook(self.files["target"])
        ws = wb[SHEET]
        ws.cell(1, COL["Column1"], "Autre chose")            # on casse une colonne obligatoire
        broken = self.folder / "CASSE.xlsx"
        wb.save(broken)
        session = Session(Selection(reference=self.files["reference"], target=broken, kam=self.files["kam"]))
        with self.assertRaises(PipelineError) as error:
            session.load()
        self.assertIn("Column1", str(error.exception))


class TestReglesLignes(BaseCase):
    def test_store_vide_ligne_entierement_rouge_et_intacte(self):
        plan, result = self.run_pipeline()
        self.assertEqual(plan.stats.red_rows, 1)
        self.assertEqual(result.colored_rows, 1)
        _wb, ws = self.result_sheet()
        for column in COL:
            self.assertEqual(ws.cell(5, COL[column]).fill.start_color.rgb, "FFFF0000",
                             f"colonne {column} non coloree")
        # aucune valeur modifiee sur cette ligne
        self.assertEqual(cell(ws, 5, "ID Promoter"), "ID_ORPHELIN")
        self.assertEqual(cell(ws, 5, "Promoter"), "PROMO ORPHELIN")
        self.assertEqual(cell(ws, 5, "Store"), None)
        self.assertEqual(cell(ws, 5, "KAM"), "KAM X")

    def test_ligne_existante_corrigee_depuis_la_reference(self):
        self.run_pipeline()
        _wb, ws = self.result_sheet()
        self.assertEqual(cell(ws, 2, "ID Promoter"), "HT204806")
        self.assertEqual(cell(ws, 2, "Promoter"), "ROUAK MOULAY HICHAM")
        self.assertEqual(cell(ws, 2, "Division"), "VD+DA")   # valeur brute de la Reference
        self.assertEqual(cell(ws, 2, "Division1"), "VD")     # division de la ligne
        self.assertEqual(cell(ws, 2, "DAY"), 26)             # Working Days, ecrit comme un nombre
        self.assertEqual(cell(ws, 2, "Column1"), "C000114373VD")

    def test_vd_plus_da_cree_une_seule_ligne_supplementaire(self):
        plan, result = self.run_pipeline()
        self.assertEqual(plan.stats.created_rows, 1)
        self.assertEqual(result.created_rows, 1)
        _wb, ws = self.result_sheet()
        self.assertEqual(ws.max_row, 10)
        self.assertEqual(cell(ws, 10, "Column1"), "C000114373DA")
        self.assertEqual(cell(ws, 10, "Division1"), "DA")
        self.assertEqual(cell(ws, 10, "ID Promoter"), "HT204806")
        self.assertEqual(cell(ws, 10, "DAY"), 26)
        # valeurs recopiees depuis la ligne soeur du meme magasin
        self.assertEqual(cell(ws, 10, "Annee"), 2026)
        self.assertEqual(cell(ws, 10, "STATUT"), "Actif")
        # IDAYA volontairement vide : aucune regle ne dit comment le calculer
        self.assertIsNone(cell(ws, 10, "IDAYA"))

    def test_pas_de_duplicate_quand_la_cle_existe(self):
        plan, _result = self.run_pipeline()
        _wb, ws = self.result_sheet()
        keys = [cell(ws, row, "Column1") for row in range(2, ws.max_row + 1)]
        self.assertEqual(len(keys), len(set(keys)), f"cles dupliquees : {keys}")
        self.assertGreaterEqual(plan.stats.duplicates_avoided, 4)

    def test_meme_store_deux_promoteurs_ne_sont_pas_melanges(self):
        self.run_pipeline()
        _wb, ws = self.result_sheet()
        self.assertEqual(cell(ws, 3, "Promoter"), "HASBI ABDELLAH")       # VD
        self.assertEqual(cell(ws, 3, "DAY"), 24)
        self.assertEqual(cell(ws, 4, "Promoter"), "BAGHDADI ABDELMOUIJIB")  # DA
        self.assertEqual(cell(ws, 4, "DAY"), 26)

    def test_kam_recupere_par_store_et_division(self):
        plan, _result = self.run_pipeline()
        _wb, ws = self.result_sheet()
        self.assertEqual(cell(ws, 2, "KAM"), "KAM VD1")
        self.assertEqual(cell(ws, 10, "KAM"), "KAM DA1")
        self.assertEqual(cell(ws, 3, "KAM"), "KAM VD2")
        self.assertEqual(plan.stats.kam_found, 5)

    def test_store_absent_de_la_reference_signale_et_intact(self):
        plan, _result = self.run_pipeline()
        self.assertEqual(plan.stats.store_not_found, 1)
        _wb, ws = self.result_sheet()
        self.assertEqual(cell(ws, 6, "ID Promoter"), "ID_INCONNU")
        self.assertEqual(cell(ws, 6, "KAM"), "KAM Y")

    def test_division_inconnue_non_traitee_mais_signalee(self):
        plan, _result = self.run_pipeline()
        messages = [entry.message for entry in plan.logger.by_level("A VERIFIER")]
        self.assertTrue(any("RAC" in message for message in messages), messages)
        _wb, ws = self.result_sheet()
        self.assertEqual(cell(ws, 7, "ID Promoter"), "RAC11111")
        self.assertEqual(cell(ws, 7, "KAM"), "KAM Z")        # ligne totalement intacte

    def test_conflit_deux_promoteurs_meme_division_non_traite(self):
        plan, _result = self.run_pipeline()
        messages = [entry.message for entry in plan.logger.by_level("A VERIFIER")]
        self.assertTrue(any("2 promoteurs differents" in message for message in messages), messages)
        _wb, ws = self.result_sheet()
        self.assertEqual(cell(ws, 8, "ID Promoter"), "CONF0000")
        self.assertEqual(cell(ws, 8, "Promoter"), "PROMO INCONNU")

    def test_cle_incoherente_rapprochee_sans_duplicate(self):
        plan, _result = self.run_pipeline()
        _wb, ws = self.result_sheet()
        self.assertEqual(cell(ws, 9, "Code Store"), "C999999999")
        self.assertEqual(cell(ws, 9, "Column1"), "C999999999VD")
        self.assertEqual(cell(ws, 9, "KAM"), "KAM VD3")
        self.assertEqual(plan.stats.created_rows, 1)         # aucune creation parasite


class TestFichiers(BaseCase):
    def test_fichiers_sources_jamais_modifies(self):
        self.run_pipeline()
        for name, path in self.files.items():
            self.assertEqual(md5(path), self.hashes[name], f"{name} a ete modifie !")

    def test_sortie_ne_peut_pas_ecraser_une_source(self):
        with self.assertRaises(PipelineError):
            self.run_pipeline(output=Path(self.files["target"]))

    def test_mise_en_forme_conservee(self):
        self.run_pipeline()
        wb, ws = self.result_sheet()
        self.assertIn("Notes", wb.sheetnames)                # les autres feuilles restent
        self.assertEqual(ws.freeze_panes, "A2")
        table = list(ws.tables.values())[0]
        self.assertEqual(table.ref, "A1:N10")                # tableau agrandi pour la ligne creee
        self.assertTrue(ws.cell(1, 1).font.b)                # en-tetes toujours en gras
        self.assertEqual(ws.cell(10, 1).font.name, ws.cell(9, 1).font.name)  # style recopie

    def test_traitement_idempotent(self):
        self.run_pipeline()
        second = self.folder / "RESULTAT2.xlsx"
        session = Session(Selection(reference=self.files["reference"], target=self.output, kam=self.files["kam"]))
        try:
            session.load()
            plan = session.analyze()
            session.apply(second, write_report=False)
        finally:
            session.close()
        self.assertEqual(plan.stats.updated_rows, 0, "un 2e passage ne doit plus rien modifier")
        self.assertEqual(plan.stats.created_rows, 0, "un 2e passage ne doit plus rien creer")

    def test_rapport_genere(self):
        session = Session(Selection(reference=self.files["reference"], target=self.files["target"],
                                    kam=self.files["kam"]))
        try:
            session.load()
            session.analyze()
            _result, reports = session.apply(self.output, write_report=True)
        finally:
            session.close()
        self.assertEqual(len(reports), 2)
        for report in reports:
            self.assertTrue(Path(report).exists())
        self.assertIn("Store vide", Path(reports[0]).read_text(encoding="utf-8"))


class TestOptions(BaseCase):
    def test_option_ajouter_les_stores_absents_du_fichier2(self):
        plan, _result = self.run_pipeline(Options(add_missing_reference_stores=True))
        _wb, ws = self.result_sheet()
        keys = [cell(ws, row, "Column1") for row in range(2, ws.max_row + 1)]
        self.assertIn("C555555555VD", keys)
        self.assertGreaterEqual(plan.stats.created_rows, 2)

    def test_option_divisions_inconnues(self):
        self.run_pipeline(Options(process_unknown_divisions=True))
        _wb, ws = self.result_sheet()
        self.assertEqual(cell(ws, 7, "ID Promoter"), "RAC11111")
        self.assertEqual(cell(ws, 7, "DAY"), 17)             # Working Days de la Reference applique
        self.assertEqual(cell(ws, 7, "Column1"), "C888888888RAC")

    def test_option_idaya_automatique(self):
        self.run_pipeline(Options(auto_number_idaya=True))
        _wb, ws = self.result_sheet()
        self.assertEqual(cell(ws, 10, "IDAYA"), 9)           # 8 lignes existantes -> 9

    def test_option_city(self):
        self.run_pipeline(Options(update_city=True))
        _wb, ws = self.result_sheet()
        self.assertEqual(cell(ws, 9, "City"), "Fes")


class TestMatcher(unittest.TestCase):
    def test_normalisation(self):
        self.assertEqual(norm_store("  Aswak   Assalam  Mohammédia "), "aswak assalam mohammedia")
        self.assertEqual(norm_store("ASWAK-ASSALAM"), "aswak assalam")
        self.assertEqual(build_key(" c003470765 ", "vd"), "C003470765VD")
        self.assertEqual(split_divisions("VD+DA"), ["VD", "DA"])
        self.assertEqual(split_divisions("VD / DA"), ["VD", "DA"])
        self.assertEqual(split_divisions(None), [])


if __name__ == "__main__":
    unittest.main(verbosity=2)
