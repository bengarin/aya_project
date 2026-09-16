"""
logger.py
---------
Le "carnet de bord" du traitement.

A chaque decision (ligne modifiee, ligne creee, ligne rouge, store introuvable...)
on ecrit une ligne dans ce carnet. A la fin, on l'affiche dans l'interface et on
peut l'exporter en .txt ou en .xlsx a cote du fichier resultat.
"""

from __future__ import annotations

import datetime as _dt
from dataclasses import dataclass, field
from pathlib import Path

# ---------------------------------------------------------------------------
# Categories d'evenements (servent aussi de filtre / couleur dans l'interface)
# ---------------------------------------------------------------------------
INFO = "INFO"
MODIF = "MODIFICATION"
CREATE = "CREATION"
RED = "LIGNE ROUGE"
STORE_NOT_FOUND = "STORE INTROUVABLE"
KAM_NOT_FOUND = "KAM INTROUVABLE"
DUPLICATE = "DUPLICATE EVITE"
CHECK = "A VERIFIER"
ERROR = "ERREUR"

ORDER = [ERROR, CHECK, RED, STORE_NOT_FOUND, KAM_NOT_FOUND, CREATE, MODIF, DUPLICATE, INFO]


@dataclass
class LogEntry:
    """Une ligne du rapport."""

    level: str
    message: str
    row: int | None = None          # numero de ligne dans le fichier 2
    store: str = ""
    division: str = ""

    def as_text(self) -> str:
        where = f"ligne {self.row}" if self.row else "-"
        store = f" [{self.store}{('/' + self.division) if self.division else ''}]" if self.store else ""
        return f"{self.level:<18} | {where:<12} |{store} {self.message}"


@dataclass
class ProcessLogger:
    """Collecte toutes les entrees du rapport, dans l'ordre du traitement."""

    entries: list[LogEntry] = field(default_factory=list)
    started_at: _dt.datetime = field(default_factory=_dt.datetime.now)

    # -- ecriture -----------------------------------------------------------
    def log(self, level: str, message: str, row: int | None = None, store: str = "", division: str = "") -> LogEntry:
        entry = LogEntry(level=level, message=message, row=row, store=store, division=division)
        self.entries.append(entry)
        return entry

    def info(self, message: str, **kw) -> LogEntry:
        return self.log(INFO, message, **kw)

    def check(self, message: str, **kw) -> LogEntry:
        """'A verifier' : situation ambigue, le logiciel ne decide PAS a ta place."""
        return self.log(CHECK, message, **kw)

    def error(self, message: str, **kw) -> LogEntry:
        return self.log(ERROR, message, **kw)

    # -- lecture ------------------------------------------------------------
    def by_level(self, level: str) -> list[LogEntry]:
        return [e for e in self.entries if e.level == level]

    def counts(self) -> dict[str, int]:
        result: dict[str, int] = {}
        for entry in self.entries:
            result[entry.level] = result.get(entry.level, 0) + 1
        return result

    def sorted_entries(self) -> list[LogEntry]:
        """Trie par importance (erreurs et 'a verifier' en premier) puis par ligne."""
        rank = {level: i for i, level in enumerate(ORDER)}
        return sorted(self.entries, key=lambda e: (rank.get(e.level, 99), e.row or 0))

    def as_text(self, header: str = "") -> str:
        lines = []
        if header:
            lines.append(header)
        lines.append(f"Rapport de traitement - {self.started_at:%d/%m/%Y %H:%M:%S}")
        lines.append("=" * 100)
        counts = self.counts()
        for level in ORDER:
            if level in counts:
                lines.append(f"{level:<18} : {counts[level]}")
        lines.append("=" * 100)
        for entry in self.sorted_entries():
            lines.append(entry.as_text())
        return "\n".join(lines)

    # -- export -------------------------------------------------------------
    def save_text(self, path: str | Path, header: str = "") -> Path:
        path = Path(path)
        path.write_text(self.as_text(header), encoding="utf-8")
        return path

    def save_xlsx(self, path: str | Path, header: str = "") -> Path:
        """Export du rapport dans un petit fichier Excel (pratique pour filtrer)."""
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        path = Path(path)
        wb = Workbook()
        ws = wb.active
        ws.title = "Rapport"
        ws.append(["Categorie", "Ligne fichier 2", "Store", "Division", "Message"])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFFFF")
            cell.fill = PatternFill("solid", start_color="FF1F4E79")
            cell.alignment = Alignment(horizontal="center")
        for entry in self.sorted_entries():
            ws.append([entry.level, entry.row, entry.store, entry.division, entry.message])
        for column, width in zip("ABCDE", (22, 16, 34, 10, 110)):
            ws.column_dimensions[column].width = width
        ws.auto_filter.ref = f"A1:E{max(ws.max_row, 1)}"
        ws.freeze_panes = "A2"
        if header:
            ws2 = wb.create_sheet("Contexte")
            for i, line in enumerate(header.splitlines(), start=1):
                ws2.cell(i, 1, line)
            ws2.column_dimensions["A"].width = 120
        wb.save(path)
        return path
